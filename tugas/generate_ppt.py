# -*- coding: utf-8 -*-
"""
Generator PowerPoint (PPTX) - Tugas Pemrograman Perangkat Bergerak
Aplikasi Pengaduan Fasilitas Kampus (Flutter + Node.js + MySQL)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image as PILImage

BASE = os.path.dirname(os.path.abspath(__file__))
MOCK = os.path.join(BASE, "mockup")
ASSETS = os.path.join(BASE, "assets")

# ── Warna tema (sama dengan aplikasi) ──
PRIMARY   = RGBColor(0x15, 0x65, 0xC0)
PRIMARY_D = RGBColor(0x0D, 0x47, 0xA1)
SECONDARY = RGBColor(0x26, 0xA6, 0x9A)
DARK      = RGBColor(0x21, 0x21, 0x21)
GREY      = RGBColor(0x60, 0x60, 0x60)
LIGHT     = RGBColor(0xF5, 0xF7, 0xFA)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
BORDER    = RGBColor(0xD0, 0xD8, 0xE0)

FONT = "Calibri"

prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ── Helper ──────────────────────────────────────────────
def add_slide():
    return prs.slides.add_slide(BLANK)

def rect(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=1):
    sp = slide.shapes.add_shape(shape, l, t, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp

def textbox(slide, l, t, w, h, text, size=18, bold=False, color=DARK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, italic=False,
            line_spacing=1.0):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size); f.bold = bold; f.italic = italic
        f.color.rgb = color; f.name = font
    return tb

def header(slide, num, title):
    """Bar header biru + nomor + judul."""
    rect(slide, 0, 0, SW, Inches(1.15), fill=PRIMARY)
    rect(slide, 0, Inches(1.15), SW, Emu(45720), fill=SECONDARY)  # garis aksen
    textbox(slide, Inches(0.55), Inches(0.18), Inches(1.2), Inches(0.8),
            num, size=36, bold=True, color=RGBColor(0x9F,0xC5,0xE8), anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, Inches(1.5), Inches(0.18), Inches(11), Inches(0.8),
            title, size=28, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

def footer(slide, n):
    textbox(slide, Inches(0.55), Inches(7.05), Inches(10), Inches(0.35),
            "Aplikasi Pengaduan Fasilitas Kampus  —  Pemrograman Perangkat Bergerak",
            size=10, color=GREY)
    textbox(slide, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.35),
            str(n), size=10, color=GREY, align=PP_ALIGN.RIGHT)

def bullets(slide, l, t, w, items, size=16, gap=8, color=DARK, bullet_color=SECONDARY):
    """items: list of (text, level) or text."""
    tb = slide.shapes.add_textbox(l, t, w, Inches(5))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    first = True
    for it in items:
        txt, lvl = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.line_spacing = 1.05
        mark = "●  " if lvl == 0 else "–  "
        r = p.add_run(); r.text = mark
        r.font.size = Pt(size); r.font.color.rgb = bullet_color if lvl==0 else GREY
        r.font.bold = True; r.font.name = FONT
        r2 = p.add_run(); r2.text = txt
        r2.font.size = Pt(size if lvl==0 else size-2)
        r2.font.color.rgb = color if lvl==0 else GREY
        r2.font.name = FONT
        if lvl > 0:
            p.level = 1
            tb.text_frame.paragraphs[-1].space_before = Pt(2)
    return tb

def add_image_fit(slide, path, cx, cy, max_w, max_h):
    """Tambah gambar, fit ke dalam kotak max_w x max_h, pusat di (cx,cy)."""
    if not os.path.exists(path):
        return None
    iw, ih = PILImage.open(path).size
    ar = iw / ih
    w = max_w; h = int(w / ar)
    if h > max_h:
        h = max_h; w = int(h * ar)
    left = int(cx - w/2); top = int(cy - h/2)
    return slide.shapes.add_picture(path, left, top, width=w, height=h)


# ══════════════════════════════════════════════════════════
#  SLIDE 1 — COVER
# ══════════════════════════════════════════════════════════
def slide_cover():
    s = add_slide()
    rect(s, 0, 0, SW, SH, fill=PRIMARY)
    rect(s, 0, 0, Inches(0.35), SH, fill=SECONDARY)
    # dekor lingkaran
    rect(s, int(SW*0.72), Inches(-1.2), Inches(4.5), Inches(4.5),
         fill=PRIMARY_D, shape=MSO_SHAPE.OVAL)
    rect(s, int(SW*0.82), Inches(4.6), Inches(3.2), Inches(3.2),
         fill=SECONDARY, shape=MSO_SHAPE.OVAL)

    # logo UMB — ditaruh di dalam lingkaran putih agar kontras di latar biru
    rect(s, Inches(0.8), Inches(0.5), Inches(1.85), Inches(1.85),
         fill=WHITE, shape=MSO_SHAPE.OVAL)
    add_image_fit(s, os.path.join(ASSETS,"logo_umb.png"),
                  Inches(1.725), Inches(1.425), Inches(1.55), Inches(1.55))

    textbox(s, Inches(1.0), Inches(2.55), Inches(9.5), Inches(0.5),
            "LAPORAN TUGAS MATA KULIAH", size=15, bold=True,
            color=RGBColor(0x9F,0xC5,0xE8))
    textbox(s, Inches(1.0), Inches(3.0), Inches(10.5), Inches(0.5),
            "PEMROGRAMAN PERANGKAT BERGERAK", size=17, bold=True,
            color=RGBColor(0xB3,0xE5,0xFC))
    textbox(s, Inches(1.0), Inches(3.75), Inches(11), Inches(1.4),
            "APLIKASI PENGADUAN\nFASILITAS KAMPUS", size=40, bold=True, color=WHITE,
            line_spacing=0.98)
    rect(s, Inches(1.05), Inches(5.15), Inches(3.2), Emu(45720), fill=SECONDARY)
    textbox(s, Inches(1.0), Inches(5.35), Inches(11), Inches(0.5),
            "Aplikasi Mobile Berbasis Client-Server untuk Pelaporan Keluhan Fasilitas Kampus",
            size=15, color=RGBColor(0xD6,0xE4,0xF7))
    textbox(s, Inches(1.0), Inches(6.15), Inches(11), Inches(0.4),
            "PROGRAM STUDI TEKNIK INFORMATIKA  •  FAKULTAS SAINS DAN TEKNOLOGI", size=12,
            color=RGBColor(0x9F,0xC5,0xE8))
    textbox(s, Inches(1.0), Inches(6.55), Inches(11), Inches(0.5),
            "UNIVERSITAS MUHAMMADIYAH BANDUNG  •  2026", size=13, bold=True,
            color=RGBColor(0x9F,0xC5,0xE8))
    return s


# ══════════════════════════════════════════════════════════
#  SLIDE 2 — ANGGOTA KELOMPOK
# ══════════════════════════════════════════════════════════
def slide_anggota():
    s = add_slide()
    header(s, "", "Anggota Kelompok")
    anggota = [
        ("Syamil Imaddudin Sinan Erhaff", "240102125"),
        ("Naufal Hafizh Wahdiyat", "240102090"),
        ("Adithya Putra Pangestu CW", "240102003"),
        ("Ahmad Ikhwal Fauzan", "240102010"),
    ]
    top = Inches(1.7)
    rowh = Inches(0.78)
    for i,(nama,nim) in enumerate(anggota):
        y = top + i*rowh
        bg = LIGHT if i % 2 == 0 else WHITE
        rect(s, Inches(2.4), y, Inches(8.5), Inches(0.66), fill=bg, line=BORDER)
        # nomor bulat
        c = rect(s, Inches(2.55), y+Inches(0.06), Inches(0.54), Inches(0.54),
                 fill=PRIMARY, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i+1)
        r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=WHITE; r.font.name=FONT
        textbox(s, Inches(3.3), y+Inches(0.05), Inches(5.5), Inches(0.55),
                nama, size=17, bold=True, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, Inches(8.8), y+Inches(0.05), Inches(2.0), Inches(0.55),
                nim, size=16, color=SECONDARY, bold=True,
                align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(2.4), Inches(5.75), Inches(8.5), Inches(0.5),
            "Program Studi Teknik Informatika — Fakultas Sains Dan Teknologi",
            size=13, color=GREY, align=PP_ALIGN.CENTER)
    footer(s, 2)
    return s


# ══════════════════════════════════════════════════════════
#  1. TOPIK
# ══════════════════════════════════════════════════════════
def slide_topik():
    s = add_slide()
    header(s, "1", "Topik")
    textbox(s, Inches(0.6), Inches(1.5), Inches(12), Inches(1.0),
            '"Perancangan dan Pembangunan Aplikasi Pengaduan Fasilitas\nKampus Berbasis Mobile dengan Arsitektur Client-Server"',
            size=20, bold=True, color=PRIMARY_D, line_spacing=1.05)
    info = [
        ("Nama Aplikasi", "Pengaduan Kampus"),
        ("Platform", "Android (Mobile) — dibangun dengan Flutter"),
        ("Arsitektur", "Client-Server (3-Tier Architecture)"),
        ("Teknologi", "Flutter (client), Node.js + Express (server), MySQL (database)"),
        ("Pengguna", "Mahasiswa (pelapor) & Admin (pengelola)"),
    ]
    top = Inches(2.75)
    rowh = Inches(0.62)
    for i,(k,v) in enumerate(info):
        y = top + i*rowh
        rect(s, Inches(0.6), y, Inches(3.0), Inches(0.52), fill=PRIMARY)
        textbox(s, Inches(0.75), y+Inches(0.03), Inches(2.8), Inches(0.46),
                k, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, Inches(3.6), y, Inches(9.1), Inches(0.52), fill=LIGHT, line=BORDER)
        textbox(s, Inches(3.8), y+Inches(0.03), Inches(8.8), Inches(0.46),
                v, size=14, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 3)
    return s


# ══════════════════════════════════════════════════════════
#  2. LATAR BELAKANG
# ══════════════════════════════════════════════════════════
def slide_latar():
    s = add_slide()
    header(s, "2", "Latar Belakang")
    items = [
        "Fasilitas kampus yang layak sangat penting untuk menunjang kegiatan belajar mengajar.",
        "Kerusakan fasilitas (lampu mati, AC, kebersihan, keamanan) sering terjadi dan mengganggu kenyamanan.",
        "Pelaporan konvensional (lisan, buku, chat) rentan hilang dan tanpa rekam jejak yang jelas.",
        "Tidak ada data terpusat untuk memantau & mengevaluasi penanganan keluhan.",
        "Hampir semua mahasiswa memiliki smartphone → aplikasi mobile jadi solusi efektif.",
        "Dibutuhkan sistem pelaporan digital yang cepat, terdokumentasi, dan terpantau statusnya.",
    ]
    bullets(s, Inches(0.7), Inches(1.6), Inches(12), items, size=17, gap=12)
    footer(s, 4)
    return s


# ══════════════════════════════════════════════════════════
#  3. RUMUSAN MASALAH
# ══════════════════════════════════════════════════════════
def slide_rumusan():
    s = add_slide()
    header(s, "3", "Rumusan Masalah")
    items = [
        "Bagaimana merancang aplikasi mobile yang memudahkan mahasiswa melaporkan keluhan fasilitas secara cepat & terdokumentasi?",
        "Bagaimana menerapkan arsitektur client-server agar aplikasi mobile dapat berkomunikasi dengan server?",
        "Bagaimana merancang autentikasi & hak akses yang membedakan peran mahasiswa dan admin?",
        "Bagaimana membangun mekanisme pelaporan berfoto dengan penelusuran status (Menunggu, Diproses, Selesai)?",
        "Bagaimana merancang basis data & API yang menyimpan data laporan secara terpusat dan terstruktur?",
    ]
    bullets(s, Inches(0.7), Inches(1.7), Inches(12), items, size=17, gap=16)
    footer(s, 5)
    return s


# ══════════════════════════════════════════════════════════
#  4. TUJUAN
# ══════════════════════════════════════════════════════════
def slide_tujuan():
    s = add_slide()
    header(s, "4", "Tujuan")
    items = [
        ("Tujuan Umum:", 0),
        ("Membangun aplikasi mobile layanan pengaduan fasilitas kampus sebagai solusi digital.", 1),
        ("Menerapkan arsitektur client-server (REST API) sehingga data tersimpan terpusat.", 1),
        ("Menyediakan autentikasi (JWT) & pembedaan peran mahasiswa dan admin.", 1),
        ("Tujuan Khusus:", 0),
        ("Memfasilitasi pelaporan berbasis bukti (kategori, lokasi, deskripsi, foto).", 1),
        ("Menyediakan penelusuran status laporan secara transparan bagi pelapor.", 1),
        ("Memenuhi tugas mata kuliah sebagai implementasi konsep koneksi server–client.", 1),
    ]
    bullets(s, Inches(0.7), Inches(1.6), Inches(12), items, size=17, gap=10)
    footer(s, 6)
    return s


# ══════════════════════════════════════════════════════════
#  5. METODE
# ══════════════════════════════════════════════════════════
def slide_metode():
    s = add_slide()
    header(s, "5", "Metode — Waterfall")
    tahap = [
        ("1", "Analisis", "Kebutuhan pengguna & fitur"),
        ("2", "Perancangan", "Arsitektur, ERD, mockup UI"),
        ("3", "Implementasi", "Backend, DB, aplikasi Flutter"),
        ("4", "Pengujian", "Uji API & alur aplikasi"),
        ("5", "Pemeliharaan", "Perbaikan & dokumentasi"),
    ]
    n = len(tahap)
    box_w = Inches(2.3); gap = Inches(0.18)
    total = n*box_w + (n-1)*gap
    left0 = int((SW - total)/2)
    y = Inches(2.0)
    for i,(num,judul,desc) in enumerate(tahap):
        x = left0 + i*(box_w+gap)
        rect(s, x, y, box_w, Inches(1.5), fill=WHITE, line=PRIMARY, line_w=1.5,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(s, x, y, box_w, Inches(0.5), fill=PRIMARY, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        textbox(s, x, y+Inches(0.02), box_w, Inches(0.46),
                f"{num}. {judul}", size=14, bold=True, color=WHITE,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x+Inches(0.12), y+Inches(0.62), box_w-Inches(0.24), Inches(0.8),
                desc, size=12, color=GREY, align=PP_ALIGN.CENTER)
        if i < n-1:
            ar = rect(s, x+box_w+Emu(22860), y+Inches(0.55), Inches(0.14), Inches(0.35),
                      fill=SECONDARY, shape=MSO_SHAPE.RIGHT_ARROW)

    # Teknologi
    textbox(s, Inches(0.7), Inches(4.0), Inches(12), Inches(0.4),
            "Teknologi yang Digunakan", size=16, bold=True, color=PRIMARY_D)
    tech = [
        ("Mobile (Client)", "Flutter (Dart)"),
        ("Server (Backend)", "Node.js + Express"),
        ("Basis Data", "MySQL / MariaDB"),
        ("Autentikasi", "JWT + Bcrypt"),
        ("Upload File", "Multer"),
        ("Pengujian API", "cURL / Postman"),
    ]
    cw = Inches(1.98); cg = Inches(0.12)
    tw = 6*cw + 5*cg
    l0 = int((SW - tw)/2)
    for i,(k,v) in enumerate(tech):
        x = l0 + i*(cw+cg)
        rect(s, x, Inches(4.55), cw, Inches(1.1), fill=LIGHT, line=BORDER,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        textbox(s, x, Inches(4.68), cw, Inches(0.4), k, size=12, bold=True,
                color=PRIMARY_D, align=PP_ALIGN.CENTER)
        textbox(s, x, Inches(5.1), cw, Inches(0.5), v, size=11, color=DARK,
                align=PP_ALIGN.CENTER)
    footer(s, 7)
    return s


# ══════════════════════════════════════════════════════════
#  6. ARSITEKTUR SISTEM
# ══════════════════════════════════════════════════════════
def slide_arsitektur():
    s = add_slide()
    header(s, "6", "Arsitektur Sistem")
    textbox(s, Inches(0.6), Inches(1.35), Inches(12), Inches(0.4),
            "Client-Server dengan pola 3-Tier Architecture", size=16, color=GREY)
    add_image_fit(s, os.path.join(MOCK,"07_arsitektur.png"),
                  int(SW/2), Inches(4.35), Inches(11.2), Inches(4.6))
    footer(s, 8)
    return s


# ══════════════════════════════════════════════════════════
#  6b. ERD
# ══════════════════════════════════════════════════════════
def slide_erd():
    s = add_slide()
    header(s, "6", "Arsitektur Sistem — Rancangan Database")
    # Tabel users
    def dbtable(x, title, rows, color):
        y = Inches(1.5)
        w = Inches(4.4)
        rect(s, x, y, w, Inches(0.5), fill=color)
        textbox(s, x, y+Inches(0.02), w, Inches(0.46), title, size=15, bold=True,
                color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        ry = y + Inches(0.5)
        for i,(f,t) in enumerate(rows):
            bg = LIGHT if i%2==0 else WHITE
            rect(s, x, ry, w, Inches(0.42), fill=bg, line=BORDER)
            textbox(s, x+Inches(0.12), ry+Inches(0.02), Inches(2.1), Inches(0.38),
                    f, size=12, bold=True, color=color, anchor=MSO_ANCHOR.MIDDLE)
            textbox(s, x+Inches(2.3), ry+Inches(0.02), Inches(2.0), Inches(0.38),
                    t, size=11, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
            ry = ry + Inches(0.42)
        return ry

    users = [("id","INT (PK)"),("nama","VARCHAR(100)"),("email","VARCHAR(100)"),
             ("password","VARCHAR(255)"),("role","ENUM"),("nim","VARCHAR(20)"),
             ("created_at","TIMESTAMP")]
    peng = [("id","INT (PK)"),("user_id","INT (FK)"),("judul","VARCHAR(150)"),
            ("kategori","VARCHAR(50)"),("lokasi","VARCHAR(150)"),("deskripsi","TEXT"),
            ("foto","VARCHAR(255)"),("status","ENUM"),("created_at","TIMESTAMP")]
    dbtable(Inches(0.8), "users", users, PRIMARY)
    dbtable(Inches(8.1), "pengaduan", peng, SECONDARY)
    # relasi
    ar = rect(s, Inches(5.5), Inches(3.1), Inches(2.4), Inches(0.45),
              fill=RGBColor(0xFF,0xF3,0xE0), line=RGBColor(0xFB,0x8C,0x00),
              shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(s, Inches(5.5), Inches(3.15), Inches(2.4), Inches(0.38),
            "1  ──<  N", size=14, bold=True, color=RGBColor(0xE6,0x51,0x00),
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.5),
            "Relasi: 1 user (mahasiswa) dapat membuat banyak laporan; setiap laporan dimiliki 1 user.",
            size=13, color=DARK, align=PP_ALIGN.CENTER)
    footer(s, 9)
    return s


# ══════════════════════════════════════════════════════════
#  7. RENCANA
# ══════════════════════════════════════════════════════════
def slide_rencana():
    s = add_slide()
    header(s, "7", "Rencana Pelaksanaan")
    rencana = [
        ("Minggu 1", "Analisis kebutuhan & studi pustaka"),
        ("Minggu 2", "Perancangan sistem (arsitektur, ERD, mockup)"),
        ("Minggu 3", "Pembangunan Backend & Database (REST API)"),
        ("Minggu 4", "Pembangunan aplikasi mobile (Flutter)"),
        ("Minggu 5", "Integrasi & pengujian end-to-end"),
        ("Minggu 6", "Penyempurnaan & dokumentasi tugas"),
    ]
    top = Inches(1.5); rowh = Inches(0.72)
    for i,(m,t) in enumerate(rencana):
        y = top + i*rowh
        rect(s, Inches(1.2), y, Inches(2.0), Inches(0.6), fill=PRIMARY)
        textbox(s, Inches(1.2), y+Inches(0.03), Inches(2.0), Inches(0.54),
                m, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE)
        rect(s, Inches(3.3), y, Inches(8.6), Inches(0.6), fill=LIGHT, line=BORDER)
        textbox(s, Inches(3.55), y+Inches(0.03), Inches(8.2), Inches(0.54),
                t, size=15, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 10)
    return s


# ══════════════════════════════════════════════════════════
#  8. MOCKUP UI MOBILE
# ══════════════════════════════════════════════════════════
def slide_mockup_grid():
    s = add_slide()
    header(s, "8", "Mockup UI Mobile")
    add_image_fit(s, os.path.join(MOCK,"09_mockup_grid.png"),
                  int(SW/2), Inches(4.35), Inches(7.6), Inches(5.4))
    footer(s, 11)
    return s

def slide_mockup_detail(pairs, page_no, title_extra=""):
    s = add_slide()
    header(s, "8", "Mockup UI Mobile" + title_extra)
    n = len(pairs)
    max_w = Inches(3.6); max_h = Inches(5.0)
    cw = int(SW / n)
    for i,(f,label) in enumerate(pairs):
        cx = int((i+0.5)*cw)
        add_image_fit(s, os.path.join(MOCK,f), cx, Inches(4.15), max_w, max_h)
        textbox(s, cx-Inches(2), Inches(6.75), Inches(4), Inches(0.4),
                label, size=15, bold=True, color=PRIMARY_D, align=PP_ALIGN.CENTER)
    footer(s, page_no)
    return s


# ══════════════════════════════════════════════════════════
#  BUILD
# ══════════════════════════════════════════════════════════
slide_cover()
slide_anggota()
slide_topik()
slide_latar()
slide_rumusan()
slide_tujuan()
slide_metode()
slide_arsitektur()
slide_erd()
slide_rencana()
slide_mockup_grid()
slide_mockup_detail([("01_splash.png","Splash Screen"),("02_login.png","Login"),("03_register.png","Register")], 12)
slide_mockup_detail([("04_home.png","Home / Daftar Laporan"),("05_buat_laporan.png","Buat Laporan"),("06_detail.png","Detail Laporan")], 13)

# slide penutup
sc = add_slide()
rect(sc, 0, 0, SW, SH, fill=PRIMARY)
rect(sc, 0, 0, Inches(0.35), SH, fill=SECONDARY)
textbox(sc, Inches(1), Inches(2.6), Inches(11), Inches(1.2),
        "TERIMA KASIH", size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rect(sc, Inches(5.4), Inches(4.0), Inches(2.5), Emu(45720), fill=SECONDARY)
textbox(sc, Inches(1), Inches(4.3), Inches(11), Inches(0.6),
        "Aplikasi Pengaduan Fasilitas Kampus", size=20, color=RGBColor(0xD6,0xE4,0xF7),
        align=PP_ALIGN.CENTER)
textbox(sc, Inches(1), Inches(4.9), Inches(11), Inches(0.6),
        "Universitas Muhammadiyah Bandung — 2026", size=16,
        color=RGBColor(0x9F,0xC5,0xE8), align=PP_ALIGN.CENTER)

out = os.path.join(BASE, "Presentasi_Tugas_Pengaduan_Kampus.pptx")
prs.save(out)
print("PPTX dibuat:", out, os.path.getsize(out), "bytes")
print("Total slide:", len(prs.slides._sldIdLst))
